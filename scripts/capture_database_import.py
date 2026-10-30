"""Capture database settings and Excel preview with disposable fictional data."""
import json
import os
from pathlib import Path
import sys
from tempfile import TemporaryDirectory

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook
from PySide6.QtCore import Qt
from PySide6.QtGui import QFont, QFontDatabase
from PySide6.QtWidgets import QApplication
from tender_app.ai_dialog import AISettingsDialog
from tender_app.database_dialog import ExcelImportDialog
from tender_app.database_excel import Issue, prepare_excel, save_template
from tender_app.database_workspace import create_empty_database
from tender_app.execution_repository import ExecutionRepository
from tender_app.execution_window import ExecutionWindow
from tender_app.settings import Settings


def main():
    output = Path(__file__).resolve().parents[1] / 'artifacts/checks/database-import/visual'
    output.mkdir(parents=True, exist_ok=True)
    app = QApplication.instance() or QApplication([])
    app.setStyle('Fusion')
    QFontDatabase.addApplicationFont('C:/Windows/Fonts/segoeui.ttf')
    app.setFont(QFont('Segoe UI', 10))
    with TemporaryDirectory(prefix='excel-visual-') as temporary:
        data = Path(temporary)
        repository = ExecutionRepository(create_empty_database(data / 'data'))
        window = ExecutionWindow(repository, Settings(data / 'data/settings.json'), data / 'documents')
        window.resize(1366, 768)
        window.show()
        settings = AISettingsDialog(window.ai_service, window)
        settings.navigation.setCurrentRow(1)
        settings.show()
        app.processEvents()
        window.grab().save(str(output / 'window-1366x768.png'))
        settings.grab().save(str(output / 'settings.png'))
        settings.changed = False
        settings.reject()
        template = output.parent / 'Шаблон_базы.xlsx'
        if not template.exists():
            save_template(template)
        book = load_workbook(template)
        for column, value in enumerate(['0001', 'Учебный завод', 'Поставка валов', '', '', '30.12.2026', '12:30', 'Победа'], 1):
            book['Тендеры'].cell(2, column, value)
        book['Спецификации'].cell(2, 1, '0001')
        book['Спецификации'].cell(2, 2, 'С1')
        book['Спецификации'].cell(2, 3, '01')
        for column, value in enumerate(['С1', 'П1', 'Вал', 'шт', 10, 8, 3], 1):
            book['Позиции'].cell(2, column, value)
        sample = output.parent / 'Вымышленный_пример.xlsx'
        book.save(sample)
        book.close()
        dialog = ExcelImportDialog(window)
        dialog.plan = prepare_excel(sample, repository)
        assert dialog.plan.valid, dialog.plan.report()
        dialog.preview.setPlainText(dialog.plan.report())
        dialog.commit_button.setEnabled(True)
        dialog.show()
        app.processEvents()
        dialog.grab().save(str(output / 'preview.png'))
        dialog.plan.issues.append(Issue(
            'Позиции', 2, 'Отгружено заказчику', 'Отгружено больше, чем размещено. Проверьте количества.'))
        dialog.preview.setPlainText(dialog.plan.report())
        dialog.commit_button.setEnabled(False)
        app.processEvents()
        dialog.grab().save(str(output / 'errors.png'))
        (output / 'result.json').write_text(json.dumps({'window': [window.width(), window.height()],
            'settings': [settings.width(), settings.height()], 'wizard': [dialog.width(), dialog.height()],
            'fictional_only': True}, ensure_ascii=False, indent=2), encoding='utf-8')
        dialog.reject()
        dialog.deleteLater()
        settings.deleteLater()
        window.dirty = False
        window.close()
        window.deleteLater()
        app.processEvents()
        repository.close()


if __name__ == '__main__':
    main()
