"""One-off import of the two execution-report layouts into a NEW database.

Run with an environment containing openpyxl. No GUI or Excel writes are used.
Private snapshots/reports belong under artifacts, never in Git or distributable seeds.
"""
import argparse
import hashlib
import json
import os
import re
import sqlite3
from contextlib import closing
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from urllib.parse import unquote

from tender_app.execution_models import Position, Specification, STAGES, UNKNOWN_UNIT, quantity
from tender_app.execution_repository import ExecutionRepository
from tender_app.models import Tender


ORDER_DUE = "Заказ покупателя.Дата отгрузки"


def text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parsed_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    for pattern in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text(value), pattern).date().isoformat()
        except ValueError:
            pass
    return None


def folder_path(link):
    if link.lower().startswith("file:///"):
        return unquote(link[8:]).replace("/", "\\")
    return ""


def read_report(source):
    from openpyxl import load_workbook

    formulas = load_workbook(source, data_only=False)
    values = load_workbook(source, data_only=True)
    specifications, loose_rows, snapshots, issues = [], [], [], []
    try:
        for ws in formulas:
            if "архив" not in ws.title.casefold() and "тендер" not in ws.title.casefold():
                continue
            cached = values[ws.title]
            headers = {text(c.value): c.column for c in ws[5] if c.value is not None}
            required = (ORDER_DUE, "Заказ покупателя.Подписан", "Заказ поставщику количество",
                        "Долг по отгрузке количество", "Номер тендера", "Статус")
            if any(name not in headers for name in required):
                raise ValueError(f"Неизвестная структура листа {ws.title!r}.")
            archive = "архив" in ws.title.casefold()
            ordered_col = 6 if archive else 7
            shipped_col = None if archive else 9
            supplier_col, document_col = (10, 11) if archive else (12, 14)
            customer, contract, active = "", "", None
            customer_raw, contract_raw = "", ""
            encountered_customer = False

            def raw_row(row):
                fields = []
                for cell in ws[row]:
                    if cell.value is None:
                        continue
                    raw = cell.value
                    if cell.data_type == "f":
                        raw = getattr(raw, "text", raw)
                        rendered = f"формула {raw}; сохранённое значение {cached.cell(row,cell.column).value!s}"
                    else:
                        rendered = text(raw)
                    if cell.hyperlink:
                        rendered += f" [ссылка: {cell.hyperlink.target or cell.hyperlink.location}]"
                    fields.append(f"{cell.coordinate} ({text(ws.cell(5, cell.column).value) or cell.column_letter}): {rendered}")
                return f"Лист «{ws.title}», строка {row}\n" + "\n".join(fields)

            def value(row, name):
                column = headers.get(name)
                return cached.cell(row, column).value if column else None

            def issue(row, message):
                issues.append({"sheet": ws.title, "row": row, "message": message})
                return message

            def stage(row, fallback, notices):
                original = text(value(row, "Статус"))
                if original in STAGES:
                    return original
                if original:
                    notices.append(issue(row, f"Исходный статус «{original}» сохранён в истории; этап требует проверки."))
                else:
                    notices.append(issue(row, "Этап не заполнен: показан этап спецификации или факт подписания."))
                return fallback

            def number(row, column, name, notices, optional=True):
                if column is None:
                    return None
                raw = cached.cell(row, column).value
                try:
                    return quantity(raw, name, optional=optional)
                except ValueError:
                    notices.append(issue(row, f"{name}: значение {raw!s} не перенесено как число; оригинал в истории."))
                    return None

            for row in range(7, ws.max_row + 1):
                name = text(cached.cell(row, 1).value)
                if not name:
                    if any(c.value is not None and c.data_type != "f" for c in ws[row]):
                        snapshots.append(raw_row(row))
                    continue
                if name.casefold() == "итого":
                    active = None
                    contract = ""
                    continue
                indent = ws.cell(row, 1).alignment.indent or 0
                raw = raw_row(row)
                snapshots.append(raw)
                if "договор" in name.casefold():
                    contract, contract_raw, active = name, raw, None
                    continue
                signed = text(value(row, "Заказ покупателя.Подписан"))
                if contract and (signed in {"Да", "Нет"} or indent in {2, 6}):
                    notices = []
                    spec_number = name
                    explicit_number = text(value(row, "Номер тендера"))
                    contacts = text(value(row, "Контакты заказчика"))
                    folder_column = headers.get("Ссылка на папку с документацией")
                    cell = ws.cell(row, folder_column)
                    link = cell.hyperlink.target if cell.hyperlink else ""
                    if not link:
                        for candidate in ws[row]:
                            if candidate.hyperlink and (candidate.hyperlink.target or "").lower().startswith("file:///"):
                                link = candidate.hyperlink.target
                                if candidate.column == headers.get("Контакты заказчика"):
                                    contacts = ""
                                break
                    path = folder_path(link or "")
                    if not explicit_number and path:
                        match = re.search(r"(?:^|[\\/])№([^_\\/\s]+)", path)
                        if match:
                            explicit_number = match[1]
                            notices.append(issue(row, "Номер тендера извлечён из названия папки по символу №, не из колонки номера."))
                    if not explicit_number:
                        notices.append(issue(row, "Номер тендера отсутствует; показана пометка «Без номера», внутренние ID сохранены раздельно."))
                    due_raw = value(row, ORDER_DUE)
                    due = parsed_date(due_raw)
                    if due_raw is not None and not due:
                        notices.append(issue(row, "Договорный срок не распознан, исходная дата сохранена в истории."))
                    eta_raw = value(row, "Ожидаемое прибытие")
                    eta = parsed_date(eta_raw)
                    if eta_raw is not None and not eta:
                        notices.append(issue(row, "Прогноз прибытия задан текстом; сохранён в истории без назначения точной даты."))
                    base_stage = stage(row, "Спецификация подписана" if signed == "Да" else "Договор на согласовании", notices)
                    document_date = re.search(r"\bот\s+(\d{2}[./]\d{2}[./]\d{4})\s*$", spec_number)
                    if "новый заказчик" in raw.casefold():
                        notices.append(issue(row, "В заметке упомянут новый заказчик: проверьте название заказчика из заголовка группы."))
                    active = dict(sheet=ws.title, row=row, customer=customer or "Заказчик не указан",
                                  tender_number=explicit_number, contacts=contacts, folder=path, link=link or "",
                                  raw=raw, customer_raw=customer_raw, contract_raw=contract_raw,
                                  notices=notices, archive=archive, source_stage=text(value(row,"Статус")),
                                  spec=Specification(number=spec_number, title=text(cached.cell(row,15).value) if not archive else "",
                                                     date=parsed_date(document_date[1]) if document_date else None,
                                                     contract=contract, signed=signed == "Да", due=due, eta=eta, stage=base_stage, archived=archive))
                    specifications.append(active)
                    encountered_customer = True
                    if archive:
                        notices.append(issue(row, "Источник — архивный лист. Запись перенесена в архив без подтверждения количеств и завершения."))
                    continue
                if indent == 0:
                    if customer and not encountered_customer:
                        loose_rows.append(dict(customer=customer, raw=customer_raw, sheet=ws.title))
                    customer, customer_raw, contract, active = name, raw, "", None
                    encountered_customer = False
                    continue
                if active is None:
                    issues.append({"sheet":ws.title,"row":row,"message":"Строка без связанной спецификации сохранена в снимке источника."})
                    continue
                notices = []
                ordered = number(row, ordered_col, "Заказано", notices, optional=False)
                if ordered is None or ordered <= 0:
                    raise ValueError(f"{ws.title}!A{row}: нельзя создать позицию без положительного заказанного количества.")
                placed = number(row, headers["Заказ поставщику количество"], "Размещено", notices)
                shipped = number(row, shipped_col, "Отгружено заказчику", notices)
                balance = number(row, headers["Долг по отгрузке количество"], "Долг", notices)
                if shipped is None and balance is not None and balance <= ordered:
                    shipped = ordered - balance
                    notices.append(issue(row, "Отгрузка заказчику рассчитана как заказано минус исходный долг по отгрузке."))
                elif shipped is not None and balance is not None and ordered-shipped != balance:
                    notices.append(issue(row, "Отгрузка и исходный долг противоречат друг другу: оставлена явная отгрузка, долг сохранён в заметке."))
                if placed is not None and (placed > ordered or (shipped is not None and shipped > placed)):
                    notices.append(issue(row, "Размещение противоречит заказу/отгрузке: оставлено неизвестным, оригинал в заметке."))
                    placed = None
                if shipped is not None and shipped > ordered:
                    notices.append(issue(row, "Отгрузка превышает заказ: оставлена неизвестной, оригинал в заметке."))
                    shipped = None
                pos_stage = stage(row, active["spec"].stage, notices)
                transport = text(value(row,"Транспорт / номер машины")) or text(value(active["row"],"Транспорт / номер машины"))
                if transport.startswith("#"):
                    notices.append(issue(row, "Ошибка формулы транспорта: в рабочее поле не перенесена, оригинал в истории."))
                    transport = ""
                location = text(value(row,"Где груз сейчас")) or text(value(active["row"],"Где груз сейчас"))
                notices.append("Единица измерения в Excel отсутствует: «ед. не указана» — пометка, а не предположение о штуках.")
                position = Position(name=name, unit=UNKNOWN_UNIT, ordered=ordered, placed=placed,
                                    customer_shipped=shipped, supplier=text(cached.cell(row,supplier_col).value),
                                    supplier_document=text(cached.cell(row,document_col).value), stage=pos_stage,
                                    transport=transport, location=location,
                                    eta=parsed_date(value(row,"Ожидаемое прибытие")) or active["spec"].eta,
                                    note=raw + "\n\nЗамечания импорта:\n" + "\n".join(notices))
                position.validate()
                active["spec"].positions.append(position)
                active["notices"].extend(notices)
            if customer and not encountered_customer:
                loose_rows.append(dict(customer=customer, raw=customer_raw, sheet=ws.title))
        if not specifications:
            raise ValueError("Спецификации в ожидаемых листах не найдены.")
        for entry in specifications:
            spec = entry["spec"]
            if not spec.title:
                spec.title = spec.positions[0].name if spec.positions else spec.number
                if len(spec.positions) > 1:
                    spec.title += f" и ещё {len(spec.positions)-1} поз."
            if not spec.positions:
                entry["notices"].append("Спецификация без позиций: требуется проверка исходного отчёта.")
            if entry["source_stage"] in {"Закрыто", "Отгружено заказчику"} and any(p.remaining != 0 for p in spec.positions):
                entry["notices"].append("Исходный статус поставки противоречит остатку: завершение не подтверждено.")
            spec.risk = "Проверить импорт: единицы измерения не указаны."
            if entry["archive"]:
                spec.risk += " Архивный лист требует проверки завершения."
            if not entry["tender_number"]:
                spec.risk += " Номер тендера отсутствует."
        return specifications, loose_rows, snapshots, issues
    finally:
        formulas.close()
        values.close()


def import_report(source, destination):
    source, destination = Path(source).resolve(), Path(destination).resolve()
    if destination.exists():
        raise FileExistsError("Импорт создаёт только новую базу; существующая не перезаписывается.")
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    entries, loose_rows, snapshots, issues = read_report(source)
    destination.parent.mkdir(parents=True, exist_ok=True)
    staging = destination.with_suffix(".importing")
    if staging.exists():
        raise FileExistsError("Найден незавершённый импорт. Выберите новый каталог.")
    schema = Path(__file__).resolve().parents[1] / "data" / "sample" / "schema.sql"
    with closing(sqlite3.connect(staging)) as connection:
        connection.executescript(schema.read_text(encoding="utf-8"))
    repo = ExecutionRepository(staging)
    tenders, mapping = {}, []
    try:
        for entry in entries:
            spec = entry["spec"]
            key = (entry["tender_number"], entry["customer"]) if entry["tender_number"] else (
                "missing", entry["sheet"], entry["row"])
            if key not in tenders:
                tender = Tender(number=entry["tender_number"] or f"Без номера · {entry['customer']}",
                                customer=entry["customer"], title=spec.title, status="Победа",
                                contacts=entry["contacts"], folder_path=entry["folder"], original_folder_link=entry["link"],
                                source_sheet=entry["sheet"], source_row=entry["row"],
                                comment=f"Импорт из {source.name}. Заказчик взят из заголовка группы Excel; при противоречии заметкам проверьте его.\n" + entry["customer_raw"],
                                import_issues="Импорт исполнения; единицы не указаны. Исходные значения сохранены в истории спецификаций.")
                tenders[key] = repo.save(tender)
            else:
                tender = tenders[key]
                changed = False
                if entry["contacts"] and entry["contacts"] not in tender.contacts:
                    tender.contacts += ("\n\n" if tender.contacts else "") + entry["contacts"]
                    changed = True
                if entry["folder"] != tender.folder_path and entry["folder"]:
                    tender.comment += "\nДругая папка спецификации: " + entry["folder"]
                    changed = True
                if changed:
                    tenders[key] = repo.save(tender)
            spec.tender_id = tenders[key].id
            note = (f"Импорт {source.name}. Договорный срок: колонка «{ORDER_DUE}» (подтверждено пользователем).\n"
                    "Отгружено означает только отгрузку заказчику. Завершение вручную не подтверждено.\n\n" +
                    entry["contract_raw"] + "\n\n" + entry["raw"] + "\n\nЗамечания импорта:\n" +
                    "\n".join(dict.fromkeys(entry["notices"])))
            saved = repo.save_specification(spec, note)
            mapping.append({"sheet":entry["sheet"],"row":entry["row"],"tender_id":saved.tender_id,
                            "specification_id":saved.id,"positions":len(saved.positions)})
        for loose in loose_rows:
            repo.save(Tender(number=f"Без номера · {loose['customer']}", customer=loose["customer"],
                             title="Отдельная строка Excel без спецификации", status="Победа", source_sheet=loose["sheet"],
                             comment=loose["raw"], import_issues="Исходная строка не содержит номера, договора или позиций; требуется проверка."))
        with repo.connection:
            for key, value in {"execution_import_sha256":source_hash,"execution_import_source":str(source),
                               "execution_import_mapping":json.dumps(mapping, ensure_ascii=False),
                               "execution_import_snapshots":json.dumps(snapshots, ensure_ascii=False)}.items():
                repo.connection.execute("INSERT INTO app_meta(key,value) VALUES (?,?)",(key,value))
        repo.check(repo.connection)
        initial_backup = repo.backup(manual=True)
        counts = {"tenders":len(repo.all()),"specifications":len(repo.specifications()),
                  "archived":sum(s.archived for s in repo.specifications()),
                  "positions":sum(len(s.positions) for s in repo.specifications()),"without_specifications":len(loose_rows)}
    finally:
        repo.close()
    if hashlib.sha256(source.read_bytes()).hexdigest() != source_hash:
        raise ValueError("Исходный Excel изменился во время импорта; база не установлена.")
    os.replace(staging, destination)
    report = {"source":str(source),"source_sha256":source_hash,"source_unchanged":True,"database":str(destination),
              "due_column":ORDER_DUE,"units":"не указаны в Excel", "completed_automatically":0,
              **counts,"mapping":mapping,"issues":issues,"backup":str(initial_backup)}
    (destination.parent / "import-report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source",type=Path)
    parser.add_argument("destination",type=Path)
    args = parser.parse_args()
    report = import_report(args.source,args.destination)
    print(json.dumps({key:report[key] for key in ("tenders","specifications","positions","without_specifications","source_unchanged")},ensure_ascii=True))


if __name__ == "__main__":
    main()
