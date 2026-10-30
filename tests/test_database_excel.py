from datetime import datetime, time
from decimal import Decimal
from pathlib import Path
import sqlite3

import pytest
from openpyxl import Workbook, load_workbook

from tender_app.database_excel import INTRO, SHEETS, prepare_excel, save_template
from tender_app.database_workspace import create_empty_database
from tender_app.execution_repository import ExecutionRepository
from tender_app.models import Tender


@pytest.fixture
def empty_repo(tmp_path):
    path = create_empty_database(tmp_path / 'База')
    repository = ExecutionRepository(path)
    yield repository
    repository.close()


def workbook(path, *, tenders=None, specs=None, positions=None, notes=None):
    book = Workbook()
    book.active.title = INTRO
    for title, data in zip(SHEETS, (tenders, specs, positions, notes)):
        sheet = book.create_sheet(title)
        sheet.append([label for _, label in SHEETS[title]])
        for record in data or []:
            sheet.append([record.get(name) for name, _ in SHEETS[title]])
    book.save(path)
    book.close()
    return path


def complete_workbook(tmp_path, **overrides):
    data = dict(
        tenders=[dict(number='001-А', customer='Завод', title='Валы', status='Победа',
                      comment='Первая строка\nВторая строка', deadline_date=datetime(2026, 12, 30),
                      deadline_time=time(12, 30), folder_path=str(tmp_path / 'Нет папки')),
                 dict(number='002', title='Без исполнения', deadline_date='31.12.2026')],
        specs=[dict(tender='001-А', code='С1', number='01', stage='В производстве', archived='Да')],
        positions=[dict(spec='С1', code='П1', name='Вал', unit='шт', ordered=10, placed=8,
                        customer_shipped=3, comment='Старый комментарий', working_note='Проверить'),
                   dict(spec='С1', code='П2', name='Втулка', unit='шт', ordered='2,5')],
        notes=[dict(kind='Тендер', code='001-А', text='Из блокнота'),
               dict(kind='Спецификация', code='С1', text='Позвонить', stamp='28.10.2026'),
               dict(kind='Позиция', code='П1', text='Ожидаем'),
               dict(kind='Позиция', code='П1', text='Получено уточнение')])
    data.update(overrides)
    return workbook(tmp_path / 'Перенос.xlsx', **data)


def test_template_is_blank_formatted_and_examples_never_imported(tmp_path):
    path = tmp_path / 'Шаблон.xlsx'
    save_template(path)
    book = load_workbook(path)
    assert book.sheetnames == [INTRO, *SHEETS]
    assert book['Тендеры']['A2'].number_format == '@'
    assert book['Тендеры'].data_validations.count
    assert all(cell.data_type != 'f' for sheet in book for row in sheet for cell in row)
    book.close()
    plan = prepare_excel(path)
    assert not plan.valid and not plan.tenders
    original = path.read_bytes()
    with pytest.raises(FileExistsError):
        save_template(path)
    assert path.read_bytes() == original


def test_full_atomic_import_preserves_data_and_comments(tmp_path, empty_repo):
    path = complete_workbook(tmp_path)
    plan = prepare_excel(path, empty_repo)
    assert plan.valid, plan.report()
    assert any(issue.warning for issue in plan.issues)
    report = empty_repo.import_excel(plan)
    assert 'Резервная копия:' in report
    assert list(empty_repo.backup_dir.glob('manual-*.db'))
    tenders = empty_repo.all()
    assert [t.number for t in tenders] == ['001-А', '002']
    assert 'Первая строка\nВторая строка\nИз блокнота' == tenders[0].comment
    assert tenders[0].deadline_time == '12:30'
    assert tenders[1].deadline_time is None
    assert tenders[0].folder_path == str(tmp_path / 'Нет папки')
    spec, = empty_repo.specifications()
    assert spec.archived
    assert spec.positions[0].remaining == Decimal(7)
    assert spec.positions[1].ordered == Decimal('2.5')
    assert spec.positions[1].placed is None and spec.positions[1].customer_shipped is None
    assert spec.positions[0].working_note == 'Проверить'
    assert len(spec.positions[0].comments) == 3
    assert spec.comments[0][1] == 'Позвонить'
    assert empty_repo.registry_comment(('specification', spec.id)) == 'Позвонить'
    assert empty_repo.registry_comment(('position', spec.id, spec.positions[0].id)) == 'Получено уточнение'
    assert spec.comments[0][0].startswith('2026-10-28')
    again = prepare_excel(path, empty_repo)
    assert again.valid and len(again.skipped) == 2
    empty_repo.import_excel(again)
    assert len(empty_repo.specifications()) == 1
    assert len(empty_repo.specifications()[0].positions[0].comments) == 3


@pytest.mark.parametrize('change', ['number', 'formula', 'unknown_column', 'unknown_sheet', 'link', 'quantity',
                                   'status', 'duplicate', 'date', 'completed', 'no_unit', 'no_ordered'])
def test_errors_block_entire_workbook(tmp_path, empty_repo, change):
    path = complete_workbook(tmp_path)
    book = load_workbook(path)
    if change == 'number': book['Тендеры']['A2'] = None
    if change == 'formula': book['Тендеры']['J2'] = '=1+1'
    if change == 'unknown_column': book['Тендеры']['Z2'] = 'Не потерять'
    if change == 'unknown_sheet': book.create_sheet('Другое')['A1'] = 'Не потерять'
    if change == 'link': book['Позиции']['A2'] = 'Нет кода'
    if change == 'quantity': book['Позиции']['G2'] = 20
    if change == 'status': book['Тендеры']['H2'] = 'Выигран'
    if change == 'duplicate': book['Тендеры']['A3'] = '001-А'
    if change == 'date': book['Тендеры']['F2'] = '31.02.2026'
    if change == 'completed': book['Спецификации']['M2'] = 'Да'
    if change == 'no_unit': book['Позиции']['D2'] = None
    if change == 'no_ordered': book['Позиции']['E2'] = None
    book.save(path)
    book.close()
    plan = prepare_excel(path, empty_repo)
    assert not plan.valid, plan.report()
    assert any(i.row and i.sheet for i in plan.issues if not i.warning)
    with pytest.raises(ValueError): empty_repo.import_excel(plan)
    assert empty_repo.all() == [] and empty_repo.specifications() == []


def test_stale_file_and_database_rejected(tmp_path, empty_repo):
    path = complete_workbook(tmp_path)
    plan = prepare_excel(path, empty_repo)
    empty_repo.save(Tender(number='999'))
    with pytest.raises(ValueError, match='изменились'): empty_repo.import_excel(plan)
    plan = prepare_excel(path, empty_repo)
    book = load_workbook(path)
    book['Тендеры']['C2'] = 'Другое'
    book.save(path)
    book.close()
    with pytest.raises(ValueError, match='изменились'): empty_repo.import_excel(plan)
    assert len(empty_repo.all()) == 1


def test_failure_after_inserts_rolls_back_whole_package(tmp_path, empty_repo):
    path = complete_workbook(tmp_path)
    empty_repo.connection.execute("CREATE TRIGGER fail_position BEFORE INSERT ON execution_positions BEGIN SELECT RAISE(ABORT, 'test failure'); END")
    empty_repo.connection.commit()
    plan = prepare_excel(path, empty_repo)
    with pytest.raises(sqlite3.IntegrityError, match='test failure'):
        empty_repo.import_excel(plan)
    assert not empty_repo.all() and not empty_repo.specifications()
    assert empty_repo.connection.execute('SELECT count(*) FROM execution_history').fetchone()[0] == 0


def test_ambiguous_existing_numbers_block(tmp_path, empty_repo):
    empty_repo.save(Tender(number='001-А'))
    empty_repo.save(Tender(number='001-А'))
    plan = prepare_excel(complete_workbook(tmp_path), empty_repo)
    assert not plan.valid and 'несколько карточек' in plan.report()


def test_existing_documents_and_resources_not_touched(tmp_path, empty_repo):
    from tender_app.resource_sync import sync_resources
    folder = tmp_path / 'Документы'
    folder.mkdir()
    document = folder / 'Договор.txt'
    document.write_text('Не менять', encoding='utf-8')
    path = workbook(tmp_path / 'Книга.xlsx', tenders=[dict(number='12345', platform='B2B', folder_path=str(folder))])
    empty_repo.import_excel(prepare_excel(path, empty_repo))
    sync_resources(empty_repo, tmp_path, lookup=lambda _: pytest.fail('Network must not run for Excel import'))
    assert list(folder.iterdir()) == [document]
    assert document.read_text(encoding='utf-8') == 'Не менять'


def test_completed_archive_and_zero_quantities(tmp_path, empty_repo):
    path = complete_workbook(tmp_path,
        specs=[dict(tender='001-А', code='С1', number='01', completed='Да', archived='Да')],
        positions=[dict(spec='С1', code='П1', name='Вал', unit='шт', ordered=2, placed=2, customer_shipped=2),
                   dict(spec='С1', code='П2', name='Втулка', unit='шт', ordered=2, placed=2, customer_shipped=2)])
    plan = prepare_excel(path, empty_repo)
    assert plan.valid, plan.report()
    empty_repo.import_excel(plan)
    assert empty_repo.specifications()[0].completed


def test_whitespace_optional_fields_and_known_zero_are_preserved(tmp_path, empty_repo):
    path = workbook(tmp_path / 'zero.xlsx',
        tenders=[dict(number='0000', status='Победа', deadline_date=' ')],
        specs=[dict(tender='0000', code='С1', number='01', stage=' ')],
        positions=[dict(spec='С1', code='П1', name='Вал', unit='шт', ordered=2,
                        placed=0, customer_shipped=0, stage=' ')])
    plan = prepare_excel(path, empty_repo)
    assert plan.valid, plan.report()
    empty_repo.import_excel(plan)
    position = empty_repo.specifications()[0].positions[0]
    assert position.placed == position.customer_shipped == Decimal(0)
    assert position.remaining == Decimal(2)


def test_change_after_backup_is_detected_under_write_lock(tmp_path, empty_repo, monkeypatch):
    path = complete_workbook(tmp_path)
    plan = prepare_excel(path, empty_repo)
    backup = empty_repo.backup
    def change_database(manual=False):
        result = backup(manual=manual)
        with empty_repo.connection:
            empty_repo.connection.execute("INSERT INTO app_meta(key,value) VALUES ('changed','yes')")
        return result
    monkeypatch.setattr(empty_repo, 'backup', change_database)
    with pytest.raises(ValueError, match='изменилась'): empty_repo.import_excel(plan)
    assert empty_repo.all() == []
