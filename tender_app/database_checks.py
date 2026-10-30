"""Disposable packaged-runtime check, invoked only by --smoke-test --check-database-import."""
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from tender_app.database_dialog import ExcelJob
from tender_app.database_excel import prepare_excel, save_template
from tender_app.database_workspace import check_workspace
from tender_app.execution_repository import ExecutionRepository


def verify_excel_runtime():
    with TemporaryDirectory(prefix='excel-runtime-') as temporary:
        root = Path(temporary)
        source = root / 'Шаблон.xlsx'
        save_template(source)
        book = load_workbook(source)
        values = {
            'Тендеры': ['0001', 'Учебный завод', 'Валы', '', '', '30.12.2026', '12:30', 'Победа'],
            'Спецификации': ['0001', 'С1', '01'],
            'Позиции': ['С1', 'П1', 'Вал', 'шт', 10, 8, 3],
            'Заметки': ['Позиция', 'П1', 'Проверка из переносимой сборки'],
        }
        for title, cells in values.items():
            for column, value in enumerate(cells, 1):
                book[title].cell(2, column, value)
        book.save(source)
        book.close()
        plan = prepare_excel(source)
        if not plan.valid:
            raise ValueError(plan.report())
        directory = root / 'employee'
        job = ExcelJob(source, None, root / 'documents', plan=plan, new_directory=directory)
        job.run()
        if job.error:
            raise ValueError(job.error)
        check_workspace(directory)
        repository = ExecutionRepository(directory / 'tenders.db')
        try:
            assert repository.all()[0].number == '0001'
            spec, = repository.specifications()
            assert len(spec.positions) == 1 and spec.positions[0].customer_shipped == 3
            assert len(spec.positions[0].comments) == 1
            duplicate = prepare_excel(source, repository)
            assert duplicate.valid and len(duplicate.skipped) == 1
            repository.import_excel(duplicate)
            assert len(repository.all()) == 1
        finally:
            repository.close()
        return True
