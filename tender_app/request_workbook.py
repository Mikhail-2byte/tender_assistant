"""Self-contained request blank, matching Tensers' workbook layout."""
from io import BytesIO
from pathlib import Path
import re

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def build_request_workbook():
    """Создаёт бланк запроса без зависимости от внешнего файла-шаблона."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.sheet_view.zoomScale = 70
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4

    widths = {
        "A": 9.14, "B": 29.43, "C": 23, "D": 20.43, "E": 26,
        "F": 14.71, "G": 14.57, "H": 13, "I": 29.57, "J": 9.14,
        "K": 13, "L": 17.86, "M": 26.71, "N": 17.14, "O": 17.57,
        "P": 13, "Q": 24.43, "R": 21.86, "S": 22.29,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    heights = {
        1: 21, 2: 29.25, 3: 21.75, 4: 21.75, 5: 89.25, 6: 21,
        7: 70.5, 8: 38.25, 18: 21, 19: 36.75, 20: 21, 21: 21,
        23: 20.25, 24: 19.5, 25: 15.75, 26: 15.75,
    }
    for row, height in heights.items():
        ws.row_dimensions[row].height = height

    thin = Side(style="thin", color="000000")
    cell_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    base_font = Font(name="Times New Roman", size=12)

    for row in ws.iter_rows(min_row=1, max_row=27, min_col=1, max_col=19):
        for cell in row:
            cell.font = base_font
            cell.alignment = Alignment(vertical="center")

    for row in range(1, 5):
        for column in range(1, 20):
            ws.cell(row, column).border = Border(
                left=thin if column == 1 else None,
                right=thin if column == 19 else None,
                top=thin,
                bottom=thin,
            )

    ws["A2"] = "Коммерческое предложение"
    ws["A2"].font = Font(name="Times New Roman", size=20)
    ws["A3"] = "Номер заявки : "
    ws["A3"].font = Font(name="Times New Roman", size=14, bold=True, italic=True)
    ws["A4"] = "Tо : "
    ws["A4"].font = Font(name="Times New Roman", size=12, bold=True, italic=True)

    headers = [
        "序号\n №", "产品名称\nНаименование", "产品材质 Материал изделия",
        "产品材料的模拟 Аналог материала изделия", "技术描述\nНомер чертежа",
        "硬度 Твёрдость, Shor", "硬度 Твёрдость, НВ", "硬度 Твёрдость, НRC",
        "技术要求 Технические требования", "单位\nЕд", "数量\nКол-во",
        "单价\nЦена ¥", "总金额\nСтоимость¥", "重量1件/公斤 Вес 1 шт/кг",
        "总重量       Общий вес, кг", "Пошлина",
        "生产厂家                                                          Завод-изготовитель",
        "生产期 Срок производства", "备注 Примечание",
    ]
    blue_fill = PatternFill("solid", fgColor="0070C0")
    red_fill = PatternFill("solid", fgColor="FF0000")
    for column, value in enumerate(headers, 1):
        cell = ws.cell(5, column, value)
        cell.font = Font(name="Times New Roman", size=12, bold=True, color="FFFFFF")
        cell.fill = blue_fill
        cell.alignment = center
        cell.border = cell_border

    sequence = {
        "A6": 1, "B6": 2, "C6": "=B6+1", "D6": "=C6+1",
        "E6": "=D6+1", "F6": "=E6+1", "G6": "=F6+1", "H6": "=G6+1",
        "I6": "=H6+1", "J6": "=I6+1", "K6": "=J6+1", "L6": "=K6+1",
        "M6": "=L6+1", "N6": "=M6+1", "O6": "=N6+1", "Q6": "=O6+1",
        "R6": "=Q6+1", "S6": "=R6+1",
    }
    for column in range(1, 20):
        cell = ws.cell(6, column)
        cell.fill = red_fill
        cell.font = Font(name="Times New Roman", size=14, color="FFFFFF")
        cell.alignment = center
        cell.border = cell_border
    for coordinate, value in sequence.items():
        ws[coordinate] = value

    for row in range(7, 9):
        for column in range(1, 20):
            ws.cell(row, column).alignment = center
            ws.cell(row, column).border = cell_border
    ws["A7"] = 1
    ws["I7"] = "По чертежу 根据图纸"
    ws["J7"] = "Штука"
    ws["M7"] = "=L7*K7"
    ws["O7"] = "=K7*N7"
    ws["S7"] = '=IFERROR(L7/N7,"")'

    ws.merge_cells("B8:L8")
    ws["B8"] = "Итого :"
    ws["B8"].alignment = Alignment(horizontal="right", vertical="center")
    ws["B8"].font = Font(name="Times New Roman", size=12, color="FF0000")
    ws["M8"] = "=SUM(M7:M7)"
    ws["O8"] = "=SUM(O7:O7)"
    ws["M8"].font = Font(name="Times New Roman", size=16, bold=True, color="FF0000")
    ws["O8"].font = Font(name="Times New Roman", size=16, bold=True, color="FF0000")

    for coordinate in ("L7", "M7", "L8", "M8"):
        ws[coordinate].number_format = '#,##0.00'
    for coordinate in ("N7", "O7", "O8", "S7"):
        ws[coordinate].number_format = '#,##0.00'

    notes = {
        "B17": "Примечание :",
        "B18": "Срок изготовления  :  …... дней  после получения аванс на заводе",
        "B19": "Условия поставки : FCA, Тяньзинь, Китай (ИНКОТЕРМС в редакции 2010 года)",
        "B20": "Срок действительна : ... дней",
        "B21": "Условия оплаты : …..% предоплата, ….% до отгрузки.",
        "E25": "00.00.2025",
    }
    for coordinate, value in notes.items():
        ws[coordinate] = value
    ws["B17"].font = Font(name="Times New Roman", size=12, color="FF0000")
    for coordinate in ("B18", "B19", "B20", "B21"):
        ws[coordinate].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    for merged_range in (
            "B18:R18", "B19:R19", "B20:E20", "B21:R21",
            "E24:R24", "B27:K27"):
        ws.merge_cells(merged_range)

    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    return wb


def create_request_workbook(folder, number):
    """Return the newly created path, or None when the request already exists."""
    safe_number = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', number).strip().rstrip('.')
    if not safe_number:
        raise ValueError('Для файла запроса нужен номер тендера.')
    target = Path(folder) / f'Запрос_{safe_number}.xlsx'
    if target.exists():
        return None
    workbook = build_request_workbook()
    try:
        with BytesIO() as buffer:
            workbook.save(buffer)
            try:
                stream = target.open('xb')
            except FileExistsError:
                return None
            try:
                with stream:
                    stream.write(buffer.getvalue())
            except BaseException:
                target.unlink(missing_ok=True)
                raise
    finally:
        workbook.close()
    return target
